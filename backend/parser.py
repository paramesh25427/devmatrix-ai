import io
from typing import Tuple, Optional
import pdfplumber
import docx

MAX_FILE_SIZE = 5 * 1024 * 1024 # 5 MB
MIN_TEXT_LENGTH = 200

class ParsingError(Exception):
    def __init__(self, code: str, message: str, status_code: int = 400):
        self.code = code
        self.message = message
        self.status_code = status_code
        super().__init__(message)

def extract_text_from_pdf(content: bytes) -> str:
    extracted_text = []
    try:
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    extracted_text.append(text)
    except Exception as e:
        raise ParsingError(
            code="no_text_found",
            message=f"Could not read PDF structure: {str(e)}. If this is a scanned document, please paste text directly.",
            status_code=422
        )
    
    full_text = "\n".join(extracted_text).strip()
    if not full_text:
        raise ParsingError(
            code="no_text_found",
            message="No readable text found in PDF. Scanned PDFs without OCR are not supported; please paste text.",
            status_code=422
        )
    return full_text

def extract_text_from_docx(content: bytes) -> str:
    try:
        doc = docx.Document(io.BytesIO(content))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        full_text = "\n".join(paragraphs).strip()
    except Exception as e:
        raise ParsingError(
            code="no_text_found",
            message=f"Could not read DOCX file: {str(e)}",
            status_code=422
        )
    
    if not full_text:
        raise ParsingError(
            code="no_text_found",
            message="No readable text found in DOCX file.",
            status_code=422
        )
    return full_text

def validate_and_extract_text(
    filename: Optional[str],
    file_bytes: Optional[bytes],
    pasted_text: Optional[str]
) -> str:
    # 1. Pasted text mode
    if pasted_text and pasted_text.strip():
        clean_text = pasted_text.strip()
        if len(clean_text) < MIN_TEXT_LENGTH:
            raise ParsingError(
                code="text_too_short",
                message=f"Requirement text must be at least {MIN_TEXT_LENGTH} characters. Provided text has {len(clean_text)} characters.",
                status_code=422
            )
        return clean_text

    # 2. File upload mode
    if not file_bytes or not filename:
        raise ParsingError(
            code="text_too_short",
            message=f"Either an uploaded file or at least {MIN_TEXT_LENGTH} characters of requirement text is required.",
            status_code=422
        )

    if len(file_bytes) > MAX_FILE_SIZE:
        raise ParsingError(
            code="file_too_large",
            message=f"Uploaded file exceeds the maximum limit of 5 MB ({len(file_bytes)} bytes).",
            status_code=400
        )

    lower_name = filename.lower()
    if lower_name.endswith(".pdf"):
        text = extract_text_from_pdf(file_bytes)
    elif lower_name.endswith(".docx") or lower_name.endswith(".doc"):
        text = extract_text_from_docx(file_bytes)
    else:
        raise ParsingError(
            code="unsupported_type",
            message=f"Unsupported file type for '{filename}'. Only PDF and DOCX files are allowed.",
            status_code=400
        )

    if len(text) < MIN_TEXT_LENGTH:
        raise ParsingError(
            code="text_too_short",
            message=f"Extracted document text is too short ({len(text)} characters). Minimum requirement is {MIN_TEXT_LENGTH} characters.",
            status_code=422
        )

    return text
